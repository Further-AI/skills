"""House style and typed-writing helpers for insurance Excel deliverables.

Every workbook built with these helpers shares one look and one set of data
rules, so files from different requests, and from follow-up turns in the same
conversation, stay consistent:

- Values keep their real types. Numbers are written as numbers with a number
  format, dates as dates, and identifiers (VINs, serials, policy and claim
  numbers) as text, so Excel never strips leading zeros or shows E-notation.
- Data sheets put the header in row 1 with nothing above it, freeze it, apply
  an autofilter over the full range, size the columns, and never merge cells,
  so every tab can be sorted, filtered, and pasted elsewhere as-is.
- The summary sheet computes counts and totals with formulas that reference the
  data sheets, so it stays correct when the user deletes resolved rows.
- Every sheet prints one page wide with the header repeated on each page, since
  these workbooks are often printed or saved as PDF for clients.

Only openpyxl is required.

Typical usage example:

  wb = new_workbook()
  columns = [
      Column("Unit #", "id"),
      Column("Year", "year"),
      Column("VIN", "id"),
      Column("Stated Value", "currency"),
  ]
  add_table_sheet(wb, "Only on Policy", columns, rows)
  only_on_policy = SummaryLine("Only on policy", count_rows("Only on Policy"))
  write_summary(
      wb["Summary"],
      title="Acme Trucking - Vehicle Schedule Reconciliation",
      sections=[Section("Results", [only_on_policy])],
  )
  wb.save("Acme_Trucking_Vehicle_Reconciliation_2026-09-11.xlsx")
"""

import datetime as dt
import json
import re
import sys
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Any, Final, Literal, Self, cast

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import (
    FormulaRule,  # pyright: ignore[reportUnknownVariableType]
)

# openpyxl's type stubs leave FormulaRule and the conditional-formatting and
# data-validation `add` methods partially untyped; the ignores below are
# limited to those call sites.
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

type ColumnKind = Literal[
    "text", "id", "integer", "year", "number", "currency", "percent", "date"
]
type CellValue = str | int | float | bool | dt.date | dt.datetime | None
type _Scalar = str | int | float | bool | dt.date

FONT_NAME: Final = "Arial"
FONT_SIZE: Final = 10
HEADER_FILL: Final = "1F3864"
STATUS_FILLS: Final = {
    "green": "E2EFDA",
    "amber": "FFF2CC",
    "red": "FCE4E4",
    "gray": "EDEDED",
    "blue": "DDEBF7",
}

_NUMBER_FORMATS: Final[dict[ColumnKind, str]] = {
    "text": "General",
    "id": "@",
    "integer": "#,##0",
    "year": "0",
    "number": "#,##0.00",
    "currency": "$#,##0;($#,##0)",
    "percent": "0.0%",
    "date": "mm/dd/yyyy",
}
# Tokens that leak from code (Python None, pandas NaN, template placeholders)
# and must never appear in a delivered cell.
_LEAKED_TOKENS: Final = frozenset(
    {"none", "nan", "nat", "null", "undefined", "<empty>", "<na>", "[object object]"}
)
# Markers people use for "no value"; in numeric and date columns they become
# blank cells so the column stays numeric.
_EN_DASH: Final = "\u2013"
_EM_DASH: Final = "\u2014"
_MISSING_MARKERS: Final = frozenset({"", "n/a", "na", "-", "--", _EN_DASH, _EM_DASH})
_DATE_FORMATS: Final = (
    "%m/%d/%Y",
    "%Y-%m-%d",
    "%m/%d/%y",
    "%m-%d-%Y",
    "%b %d, %Y",
    "%B %d, %Y",
    "%d-%b-%Y",
)
_NUMBER_TEXT: Final = re.compile(r"^\(?-?\$?\s*-?[\d,]*\.?\d+\s*\)?$")
_INVALID_TITLE_CHARS: Final = re.compile(r"[\[\]:*?/\\]")
_MAX_TITLE_LENGTH: Final = 31
_MIN_WIDTH: Final = 8
_MAX_WIDTH: Final = 50
_MAX_WRAP_WIDTH: Final = 60
_DROPDOWN_EXTRA_ROWS: Final = 500
_PORTRAIT_MAX_COLUMNS: Final = 5
# Custom document property listing comparison-matrix sheets, so
# check_workbook.py can skip list-only rules (autofilter, mixed types) there.
MATRIX_SHEETS_PROPERTY: Final = "excel-generation:matrix-sheets"


class Formula(str):
    """Marks an intentionally authored Excel formula, never extracted source text."""

    def __new__(cls, expression: str) -> Self:
        """Creates a formula from an expression beginning with '='.

        Args:
            expression: Trusted formula including its leading equals sign.

        Returns:
            A formula that workbook writers may execute.

        Raises:
            ValueError: If the expression lacks '=' or a formula body.
        """
        if not expression.startswith("=") or not expression[1:].strip():
            raise ValueError("A formula must start with '=' and have a body")
        return super().__new__(cls, expression)


def _write_cell(ws: Worksheet, *, row: int, column: int, value: CellValue) -> Cell:
    """Writes literal strings as text; only explicit Formula values execute."""
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and not isinstance(value, Formula):
        # openpyxl otherwise interprets source strings beginning with '=' as formulas.
        cell.data_type = "s"
    return cell


@dataclass(frozen=True, slots=True)
class Column:
    """One column of a data sheet.

    Attributes:
        header: Header text shown in row 1. Name sources by their business role
            ("Policy Year", "Fleet List Year"), not by file type ("PDF Year").
        kind: How values are coerced and formatted. Use "id" for anything that
            identifies rather than measures: VINs, serials, unit, policy, claim,
            and location numbers, ZIP codes.
        width: Column width in characters. None auto-fits to the content.
        number_format: Excel number format overriding the kind's default, for
            example "$#,##0.00" for premiums with cents.
        wrap: Whether long text wraps within the cell instead of widening the
            column.
    """

    header: str
    kind: ColumnKind = "text"
    width: float | None = None
    number_format: str | None = None
    wrap: bool = False


@dataclass(frozen=True, slots=True)
class SummaryLine:
    """One labeled value on the summary sheet.

    Attributes:
        label: What the value is, in the reader's terms.
        value: A literal value or an explicit `Formula` for a calculation.
        note: Optional explanation shown in the third column.
        number_format: Excel number format for the value. Integers and formulas
            default to "#,##0"; pass a currency or percent format when needed.
    """

    label: str
    value: CellValue = None
    note: str | None = None
    number_format: str | None = None


@dataclass(frozen=True, slots=True)
class Section:
    """A titled block of lines on the summary sheet.

    Attributes:
        heading: Bold heading above the block, or None for no heading.
        lines: The lines in display order.
        column_headers: Optional header row for the block, for example
            ("Category", "Count", "What it means").
    """

    heading: str | None
    lines: Sequence[SummaryLine] = ()
    column_headers: tuple[str, str, str] | None = None


@dataclass(frozen=True, slots=True)
class MatrixRow:
    """One row of a comparison matrix.

    Attributes:
        label: Row label shown in column A, e.g. "Each Occurrence".
        values: One value per compared column, in column order. Text such as
            "Included" or "Not quoted" is kept as-is.
        kind: How numeric values in this row are coerced and formatted.
        note: Optional note shown in the last column.
        number_format: Excel number format overriding the kind's default, for
            example "$#,##0.00" for premiums.
    """

    label: str
    values: Sequence[object] = ()
    kind: ColumnKind = "currency"
    note: str | None = None
    number_format: str | None = None


@dataclass(frozen=True, slots=True)
class MatrixGroup:
    """A titled group of matrix rows, such as one line of coverage.

    Attributes:
        heading: Group title shown on its own shaded row, e.g. "Auto".
        rows: The group's rows in display order.
    """

    heading: str
    rows: Sequence[MatrixRow] = ()


def new_workbook(summary_title: str = "Summary") -> Workbook:
    """Creates an empty workbook whose first sheet is the summary sheet.

    Args:
        summary_title: Title of the first sheet.

    Returns:
        A workbook with one empty sheet named `summary_title`.
    """
    wb = Workbook()
    summary = wb.active
    if summary is None:  # Unreachable for a new Workbook; narrows the type.
        summary = wb.create_sheet()
    summary.title = safe_sheet_title(summary_title)
    return wb


def safe_sheet_title(title: str, existing: Sequence[str] = ()) -> str:
    """Returns a valid, portable, unique Excel sheet title.

    Removes characters Excel forbids, drops emoji and other non-ASCII symbols
    (they break formula references and some import tools), replaces en and em
    dashes with hyphens, truncates to 31 characters, and appends a counter if
    the title is already taken.

    Args:
        title: Desired title.
        existing: Titles already in the workbook, compared case-insensitively.
    """
    cleaned = title.replace(_EN_DASH, "-").replace(_EM_DASH, "-")
    cleaned = _INVALID_TITLE_CHARS.sub(" ", cleaned)
    cleaned = cleaned.encode("ascii", "ignore").decode()
    cleaned = " ".join(cleaned.split())[:_MAX_TITLE_LENGTH].strip() or "Sheet"
    taken = {name.lower() for name in existing}
    candidate, counter = cleaned, 2
    while candidate.lower() in taken:
        suffix = f" ({counter})"
        candidate = f"{cleaned[: _MAX_TITLE_LENGTH - len(suffix)]}{suffix}"
        counter += 1
    return candidate


def coerce(value: object, kind: ColumnKind) -> CellValue:
    """Converts a raw extracted value to the type its column requires.

    Leaked code tokens ("None", "nan", "<empty>") become blank in every kind.
    In numeric and date kinds, "no value" markers such as "N/A" and "—" also
    become blank. A value that can't be converted is returned unchanged, so no
    data is lost; `add_table_sheet` reports it so it can be fixed.

    Args:
        value: Raw value from a parsed document, spreadsheet, or computation.
        kind: Target column kind. Percent strings like "94%" become 0.94;
            numbers passed for a percent column must already be fractions.

    Returns:
        The converted value, None for blanks, or the original value if it
        could not be converted.
    """
    if value is None or isinstance(value, Formula):
        return value
    if not isinstance(value, str | int | float | bool | dt.date):
        value = str(value)  # Decimals, numpy scalars, and the like.
    if isinstance(value, str):
        text = value.strip()
        if text.lower() in _LEAKED_TOKENS:
            return None
        value = text
    if kind == "text":
        return None if value == "" else value
    if kind == "id":
        return _to_id(value)
    if isinstance(value, str) and value.lower() in _MISSING_MARKERS:
        return None
    match kind:
        case "integer" | "year":
            return _to_integer(value, is_year=kind == "year")
        case "number" | "currency":
            return _to_float(value)
        case "percent":
            return _to_percent(value)
        case "date":
            return _to_date(value)


def add_table_sheet(
    wb: Workbook,
    title: str,
    columns: Sequence[Column],
    rows: Sequence[Mapping[str, object] | Sequence[object]],
    *,
    freeze_first_column: bool = False,
) -> Worksheet:
    """Adds a data sheet in the house style.

    Writes the header in row 1 and one record per row below it, with values
    coerced to each column's kind. Freezes the header, applies an autofilter
    over the full range, and sizes the columns. Values that could not be
    converted are written as-is and listed on stderr so they can be fixed.

    Args:
        wb: Workbook to add the sheet to.
        title: Sheet title; made valid and unique with `safe_sheet_title`.
        columns: Column definitions in display order.
        rows: Records, each either a mapping keyed by column header or a
            sequence with exactly one value per column. Missing mapping keys
            become blank cells.
        freeze_first_column: Also freeze column A, useful for wide sheets
            whose first column identifies the record.

    Returns:
        The new worksheet.

    Raises:
        ValueError: If a sequence row does not match the column count.
    """
    ws = wb.create_sheet(safe_sheet_title(title, wb.sheetnames))
    headers = [column.header for column in columns]
    for col, header in enumerate(headers, start=1):
        _write_cell(ws, row=1, column=col, value=header)
    _style_header_row(ws, len(columns))

    problems: list[str] = []
    for row_number, record in enumerate(rows, start=2):
        raw_values = (
            [record.get(header) for header in headers]
            if isinstance(record, Mapping)
            else list(record)
        )
        if len(raw_values) != len(columns):
            raise ValueError(
                f"Row {row_number} has {len(raw_values)} values; expected {len(columns)}"
            )
        for col_number, (column, raw) in enumerate(
            zip(columns, raw_values, strict=True), 1
        ):
            value = coerce(raw, column.kind)
            cell = _write_cell(ws, row=row_number, column=col_number, value=value)
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE)
            cell.number_format = column.number_format or _NUMBER_FORMATS[column.kind]
            if column.wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if _is_unconverted(value, column.kind):
                problems.append(
                    f"{ws.title}!{cell.coordinate} ({column.header}): {value!r}"
                )

    last_row = max(ws.max_row, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"
    ws.freeze_panes = "B2" if freeze_first_column else "A2"
    _size_columns(ws, columns)
    _set_print_layout(ws, landscape=len(columns) > _PORTRAIT_MAX_COLUMNS)
    if problems:
        print(
            f"[xlsx_kit] {len(problems)} value(s) kept as text because they don't fit "
            f"their column kind:\n  " + "\n  ".join(problems[:20]),
            file=sys.stderr,
        )
    return ws


def add_matrix_sheet(
    wb: Workbook,
    title: str,
    *,
    row_header: str,
    columns: Sequence[str],
    groups: Sequence[MatrixGroup],
    baseline: str | None = None,
    notes_header: str | None = "Notes",
) -> Worksheet:
    """Adds a comparison matrix: items down the side, options across the top.

    For quote, renewal, and coverage comparisons, where each row has its own
    type (a limit, a percentage, a premium) and cells may hold text such as
    "Excluded". Group headings get their own shaded row. When `baseline` names
    one of `columns` (usually "Expiring"), other columns' cells that differ
    from it are shaded amber, and the shading follows later edits.

    Args:
        wb: Workbook to add the sheet to.
        title: Sheet title; made valid and unique with `safe_sheet_title`.
        row_header: Header of the label column, e.g. "Coverage".
        columns: Headers of the compared columns, e.g. carrier names.
        groups: Row groups in display order.
        baseline: Column to compare the others against, or None.
        notes_header: Header of a trailing notes column, or None to omit it.

    Returns:
        The new worksheet.

    Raises:
        ValueError: If `baseline` is not one of `columns`, or a row has more
            values than there are columns.
    """
    if baseline is not None and baseline not in columns:
        raise ValueError(
            f"Baseline {baseline!r} is not one of the columns {list(columns)}"
        )
    ws = wb.create_sheet(safe_sheet_title(title, wb.sheetnames))
    headers = [row_header, *columns, *([notes_header] if notes_header else [])]
    for col, header in enumerate(headers, start=1):
        _write_cell(ws, row=1, column=col, value=header)
    _style_header_row(ws, len(headers))
    for col in range(2, len(columns) + 2):  # Align headers with their numbers.
        ws.cell(row=1, column=col).alignment = Alignment(
            horizontal="right", vertical="center", wrap_text=True
        )

    group_fill = PatternFill(
        "solid", start_color=STATUS_FILLS["blue"], end_color=STATUS_FILLS["blue"]
    )
    row_number = 2
    for group in groups:
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_number, column=col)
            cell.fill = group_fill
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
        _write_cell(ws, row=row_number, column=1, value=group.heading)
        row_number += 1
        for matrix_row in group.rows:
            if len(matrix_row.values) > len(columns):
                raise ValueError(
                    f"Row {matrix_row.label!r} has more values than columns"
                )
            _write_matrix_row(
                ws,
                row_number,
                matrix_row,
                notes_column=len(columns) + 2 if notes_header else None,
            )
            row_number += 1

    if baseline is not None:
        _shade_differences(ws, columns, baseline, last_row=row_number - 1)
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 36
    for col in range(2, len(columns) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 18
    if notes_header:
        ws.column_dimensions[get_column_letter(len(headers))].width = 50
    _set_print_layout(ws, landscape=len(headers) > _PORTRAIT_MAX_COLUMNS)
    _register_matrix_sheet(wb, ws.title)
    return ws


def color_rows(ws: Worksheet, header: str, colors: Mapping[str, str]) -> None:
    """Shades whole data rows based on the value in one column.

    Uses conditional formatting, so the shading follows the value if the user
    edits it later (for example changing "Review" to "Resolved").

    Args:
        ws: A sheet written by `add_table_sheet`.
        header: Header of the column whose value drives the color.
        colors: Maps an exact cell value to a color name from `STATUS_FILLS`
            ("green", "amber", "red", "gray", "blue") or a hex RGB string.
    """
    letter = column_letter(ws, header)
    last_col = get_column_letter(ws.max_column)
    data_range = f"A2:{last_col}{max(ws.max_row, 2)}"
    for value, color in colors.items():
        fill_color = STATUS_FILLS.get(color, color)
        escaped = value.replace('"', '""')
        rule = FormulaRule(  # pyright: ignore[reportUnknownVariableType]
            formula=[f'${letter}2="{escaped}"'],
            fill=PatternFill("solid", start_color=fill_color, end_color=fill_color),
        )
        ws.conditional_formatting.add(data_range, rule)  # pyright: ignore[reportUnknownMemberType, reportUnknownArgumentType]


def add_dropdown(ws: Worksheet, header: str, options: Sequence[str]) -> None:
    """Restricts a column to a list of choices, shown as an in-cell dropdown.

    Useful for an "Action" or "Resolution" column on a worklist tab. The
    dropdown extends 500 rows past the data so rows the user adds get it too.

    Args:
        ws: A sheet written by `add_table_sheet`.
        header: Header of the column to restrict.
        options: Allowed values. Their combined length must stay under 255
            characters, an Excel limit for inline lists.
    """
    letter = column_letter(ws, header)
    formula = '"' + ",".join(option.replace('"', "") for option in options) + '"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    cells = f"{letter}2:{letter}{ws.max_row + _DROPDOWN_EXTRA_ROWS}"
    validation.add(cells)  # pyright: ignore[reportUnknownMemberType]
    ws.add_data_validation(validation)


def column_letter(ws: Worksheet, header: str) -> str:
    """Returns the column letter of a header in row 1.

    Args:
        ws: Sheet whose header row to search.
        header: Exact header text.

    Raises:
        KeyError: If no cell in row 1 contains exactly `header`.
    """
    for index, cell in enumerate(ws[1], start=1):
        if cell.value == header:
            return get_column_letter(index)
    raise KeyError(f"No column headed {header!r} on sheet {ws.title!r}")


def count_rows(sheet_title: str, column: str = "A") -> Formula:
    """Returns a formula counting the data rows of a sheet.

    Counts non-blank cells below the header, so pick a column that is filled
    on every row (the record key or source reference, usually column A).

    Args:
        sheet_title: Title of the data sheet, exactly as created.
        column: Column letter to count.
    """
    return Formula(f"=COUNTA({quote_sheetname(sheet_title)}!{column}2:{column}1048576)")


def sum_column(sheet_title: str, column: str) -> Formula:
    """Returns a formula summing a numeric column of a data sheet.

    Args:
        sheet_title: Title of the data sheet, exactly as created.
        column: Column letter to sum; get it with `column_letter`.
    """
    return Formula(f"=SUM({quote_sheetname(sheet_title)}!{column}2:{column}1048576)")


def balance_check(total: str | int, parts: Sequence[str | int]) -> Formula:
    """Returns a formula showing "OK" when a total equals the sum of its parts.

    Use it on the summary sheet to prove every source record landed in exactly
    one bucket, for example policy vehicles = matched + needs review + only on
    policy.

    Args:
        total: Literal count, or a formula such as `count_rows("Policy Extract")`.
        parts: Literal counts or formulas, such as `count_rows("Matched")`.

    Returns:
        A formula like `=IF(94=COUNTA(...)+COUNTA(...),"OK","MISMATCH")`.
    """

    def term(value: str | int) -> str:
        return value.removeprefix("=") if isinstance(value, str) else str(value)

    summed = "+".join(term(part) for part in parts) or "0"
    return Formula(f'=IF({term(total)}={summed},"OK","MISMATCH")')


def write_summary(
    ws: Worksheet,
    *,
    title: str,
    sections: Sequence[Section],
    subtitle: str | None = None,
) -> None:
    """Writes the summary sheet: a title block followed by labeled sections.

    Typical sections, in order: key facts (insured, policy, period, carrier),
    results with counts as formulas, reconciliation checks, sources, and
    method and assumptions.

    Args:
        ws: The (empty) summary sheet.
        title: Bold title in A1, for example "Acme Trucking - Vehicle Audit".
        sections: Blocks written top to bottom with a blank row between them.
        subtitle: Optional line under the title, for example the preparation
            date and source documents.
    """
    _write_cell(ws, row=1, column=1, value=title)
    ws["A1"].font = Font(name=FONT_NAME, size=14, bold=True, color=HEADER_FILL)
    row = 2
    if subtitle:
        _write_cell(ws, row=row, column=1, value=subtitle).font = Font(
            name=FONT_NAME, size=FONT_SIZE, italic=True, color="595959"
        )
        row += 1

    for section in sections:
        row += 1  # Blank row before every section.
        if section.heading:
            _write_cell(ws, row=row, column=1, value=section.heading).font = Font(
                name=FONT_NAME, size=11, bold=True, color=HEADER_FILL
            )
            row += 1
        if section.column_headers:
            for col, text in enumerate(section.column_headers, start=1):
                _write_cell(ws, row=row, column=col, value=text)
            _style_header_row(ws, len(section.column_headers), row=row)
            row += 1
        for line in section.lines:
            _write_summary_line(ws, row, line)
            row += 1

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 70
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _write_matrix_row(
    ws: Worksheet, row: int, matrix_row: MatrixRow, *, notes_column: int | None
) -> None:
    """Writes one matrix row: label, typed values, and an optional note."""
    font = Font(name=FONT_NAME, size=FONT_SIZE)
    label_cell = _write_cell(ws, row=row, column=1, value=matrix_row.label)
    label_cell.font = font
    label_cell.alignment = Alignment(indent=1, vertical="top")
    for offset, raw in enumerate(matrix_row.values):
        cell = _write_cell(
            ws, row=row, column=2 + offset, value=coerce(raw, matrix_row.kind)
        )
        cell.font = font
        cell.number_format = (
            matrix_row.number_format or _NUMBER_FORMATS[matrix_row.kind]
        )
        cell.alignment = Alignment(horizontal="right", wrap_text=True, vertical="top")
    if notes_column is not None and matrix_row.note:
        note_cell = _write_cell(ws, row=row, column=notes_column, value=matrix_row.note)
        note_cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color="595959")
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")


def _shade_differences(
    ws: Worksheet, columns: Sequence[str], baseline: str, *, last_row: int
) -> None:
    """Shades matrix cells that differ from the baseline column."""
    base_letter = get_column_letter(2 + list(columns).index(baseline))
    fill = PatternFill(
        "solid", start_color=STATUS_FILLS["amber"], end_color=STATUS_FILLS["amber"]
    )
    for offset, header in enumerate(columns):
        if header == baseline:
            continue
        letter = get_column_letter(2 + offset)
        rule = FormulaRule(  # pyright: ignore[reportUnknownVariableType]
            formula=[
                f'AND({letter}2<>"",${base_letter}2<>"",{letter}2<>${base_letter}2)'
            ],
            fill=fill,
        )
        ws.conditional_formatting.add(f"{letter}2:{letter}{max(last_row, 2)}", rule)  # pyright: ignore[reportUnknownMemberType, reportUnknownArgumentType]


def _register_matrix_sheet(wb: Workbook, title: str) -> None:
    """Records a matrix sheet's title in the workbook's custom properties."""
    # openpyxl's type stubs predate custom document properties, so this one
    # access is treated as Any; the values stored are plain strings.
    props = cast(Any, wb).custom_doc_props
    existing = next(
        (prop for prop in props.props if prop.name == MATRIX_SHEETS_PROPERTY), None
    )
    titles: list[str] = json.loads(str(existing.value)) if existing is not None else []
    titles.append(title)
    if existing is not None:
        props.props.remove(existing)
    props.append(StringProperty(name=MATRIX_SHEETS_PROPERTY, value=json.dumps(titles)))


def _write_summary_line(ws: Worksheet, row: int, line: SummaryLine) -> None:
    """Writes one summary line across columns A to C."""
    font = Font(name=FONT_NAME, size=FONT_SIZE)
    _write_cell(ws, row=row, column=1, value=line.label).font = font
    value_cell = _write_cell(ws, row=row, column=2, value=line.value)
    value_cell.font = font
    value_cell.alignment = Alignment(horizontal="right")
    is_formula = isinstance(line.value, Formula)
    if line.number_format:
        value_cell.number_format = line.number_format
    elif is_formula or isinstance(line.value, int):
        value_cell.number_format = "#,##0"
    elif isinstance(line.value, str):
        # Long text values wrap instead of running under the note column.
        value_cell.alignment = Alignment(
            horizontal="left", wrap_text=True, vertical="top"
        )
    if line.note:
        note_cell = _write_cell(ws, row=row, column=3, value=line.note)
        note_cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color="595959")
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")


def _set_print_layout(ws: Worksheet, *, landscape: bool) -> None:
    """Prints the sheet one page wide with the header row repeated."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0  # As many pages tall as needed.
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"


def _style_header_row(ws: Worksheet, width: int, row: int = 1) -> None:
    """Applies the header style to the first `width` cells of a row."""
    fill = PatternFill("solid", start_color=HEADER_FILL, end_color=HEADER_FILL)
    for col in range(1, width + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _size_columns(ws: Worksheet, columns: Sequence[Column]) -> None:
    """Sets each column's width from its definition or its content."""
    sample_rows = min(ws.max_row, 500)
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        if column.width is not None:
            ws.column_dimensions[letter].width = column.width
            continue
        longest = len(column.header)
        for (value,) in ws.iter_rows(
            min_row=2,
            max_row=sample_rows,
            min_col=index,
            max_col=index,
            values_only=True,
        ):
            longest = max(longest, _display_length(value))
        limit = _MAX_WRAP_WIDTH if column.wrap else _MAX_WIDTH
        ws.column_dimensions[letter].width = min(max(longest + 2, _MIN_WIDTH), limit)


def _display_length(value: object) -> int:
    """Estimates how many characters a value occupies once formatted."""
    if value is None:
        return 0
    if isinstance(value, dt.date):
        return 10
    if isinstance(value, float):
        return len(f"{value:,.2f}") + 1
    if isinstance(value, int):
        return len(f"{value:,}") + 1
    return len(str(value))


def _is_unconverted(value: object, kind: ColumnKind) -> bool:
    """Returns whether a coerced value still doesn't match its column kind."""
    if value is None or kind in {"text", "id"}:
        return False
    if isinstance(value, Formula):
        return False
    return isinstance(value, str)


def _to_id(value: _Scalar) -> str | None:
    """Converts an identifier to text, undoing float artifacts like '123.0'."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _to_integer(value: _Scalar, *, is_year: bool) -> CellValue:
    """Converts to int, or returns the value unchanged if it isn't integral."""
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        digits = value.replace(",", "")
        if digits.lstrip("-").isdigit():
            number = int(digits)
            if not is_year or 1800 <= number <= 2200:
                return number
    return value


def _to_float(value: _Scalar) -> CellValue:
    """Converts numbers and currency text such as '$1,250' or '(300)'."""
    if isinstance(value, bool):
        return value
    if isinstance(value, int | float):
        return value
    if isinstance(value, str) and _NUMBER_TEXT.match(value):
        negative = value.startswith("(") and value.endswith(")")
        cleaned = value.strip("()").replace("$", "").replace(",", "").replace(" ", "")
        number = float(cleaned)
        number = -number if negative else number
        return int(number) if number.is_integer() else number
    return value


def _to_percent(value: _Scalar) -> CellValue:
    """Converts '94%' to 0.94; numbers are assumed to be fractions already."""
    if isinstance(value, str) and value.endswith("%"):
        number = _to_float(value[:-1].strip())
        if isinstance(number, int | float):
            return number / 100
        return value
    return _to_float(value)


def _to_date(value: _Scalar) -> CellValue:
    """Converts datetimes and common US and ISO date strings to dates."""
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        for date_format in _DATE_FORMATS:
            try:
                return dt.datetime.strptime(value, date_format).date()  # noqa: DTZ007 — calendar date, not an instant
            except ValueError:
                continue
    return value
